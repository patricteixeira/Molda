"""Constrói os templates OOXML versionados usados pelo worker do M2."""

from __future__ import annotations

import argparse
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Mm
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DESTINATION = ROOT / "apps" / "api" / "src" / "brand_api" / "native_templates" / "v1"
FIXED_TIME = datetime(2026, 7, 14, 12, 0, 0)
FIXED_ZIP_TIME = (2026, 7, 14, 12, 0, 0)
PPTX_PROFILES = {
    "post-1x1": (10.0, 10.0),
    "post-4x5": (8.0, 10.0),
    "story-9x16": (7.5, 13.333333),
}


def _normalize_package(path: Path) -> None:
    """Reescreve o ZIP em ordem estável e remove timestamps do ambiente."""
    with zipfile.ZipFile(path) as source:
        payloads = {
            info.filename: (source.read(info.filename), info.is_dir())
            for info in source.infolist()
        }
    temporary = path.with_suffix(path.suffix + ".tmp")
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name in sorted(payloads):
                payload, is_directory = payloads[name]
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.create_system = 3
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = (0o40755 if is_directory else 0o100644) << 16
                target.writestr(info, payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _set_core_properties(properties) -> None:
    properties.author = "Brand Runtime"
    properties.last_modified_by = "Brand Runtime"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    properties.title = "Brand Runtime native template v1"
    properties.subject = "Versioned template-fill source"


def _build_pptx(path: Path, width_inches: float, height_inches: float) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(width_inches)
    presentation.slide_height = Inches(height_inches)
    _set_core_properties(presentation.core_properties)
    compatible = next(
        (layout for layout in presentation.slide_layouts if layout.name == "Title and Content"),
        None,
    )
    if compatible is None:
        raise RuntimeError("O template-base não contém o layout Title and Content.")
    presentation.save(path)
    _normalize_package(path)


def _build_docx(path: Path) -> None:
    document = Document()
    _set_core_properties(document.core_properties)
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.right_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    for slot_id in ("title", "body", "logo"):
        document.add_paragraph(f"{{{{slot:{slot_id}}}}}")
    document.save(path)
    _normalize_package(path)


def build(destination: Path) -> list[Path]:
    """Materializa os quatro recursos de template da versão atual."""
    destination.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    for profile, (width, height) in PPTX_PROFILES.items():
        path = destination / f"pptx-{profile}.pptx"
        _build_pptx(path, width, height)
        generated.append(path)
    docx = destination / "docx-doc-a4.docx"
    _build_docx(docx)
    generated.append(docx)
    return generated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--destination", type=Path, default=DEFAULT_DESTINATION)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if not args.check:
        for path in build(args.destination):
            print(path.relative_to(ROOT))
        return

    with tempfile.TemporaryDirectory(prefix="brandrt-native-templates-") as raw_temp:
        rebuilt = build(Path(raw_temp))
        mismatches = [
            path.name
            for path in rebuilt
            if not (args.destination / path.name).is_file()
            or path.read_bytes() != (args.destination / path.name).read_bytes()
        ]
    if mismatches:
        raise SystemExit("Templates desatualizados: " + ", ".join(mismatches))
    print("Templates nativos v1 estão reproduzíveis.")


if __name__ == "__main__":
    main()
